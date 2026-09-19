#!/usr/bin/env python3
"""Canonical QA workbook for Sessions Dev cloud-analysis run QA-0001."""
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent / "session-insights-qa.xlsx"
NOW = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
BUILD = "Sessions Dev.app sha256:80a93ed5d29c1afa mtime 2026-09-19T14:49Z"
ENV = "me.cepessa.sessions-dev; CEPESSA_INSIGHTS_USE_FIXTURE=1; test-root=/private/tmp/cepessa-sessions-jev/ui-root"
EV = "qa/evidence"

header_fill = PatternFill("solid", fgColor="1F2937")
header_font = Font(color="FFFFFF", bold=True, name="SF Pro Text", size=11)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def style_header(ws, n):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n)}1"
    for col in range(1, n + 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.row_dimensions[1].height = 28


def add_rows(ws, headers, rows):
    ws.append(headers)
    style_header(ws, len(headers))
    for row in rows:
        ws.append(list(row))
        r = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.alignment = wrap
            cell.border = thin
            cell.font = Font(name="SF Pro Text", size=10)
        ws.row_dimensions[r].height = 48


def main():
    wb = Workbook()

    # --- Runs ---
    ws = wb.active
    ws.title = "Runs"
    add_rows(
        ws,
        [
            "run_id", "app", "repository", "commit_or_build", "artifact_path",
            "platform", "environment", "roles", "permissions", "locale",
            "feature_flags", "integrations", "scope", "exclusions",
            "protected_actions", "initial_token_budget", "remaining_tokens_checkpoint",
            "phase", "status", "started_at", "updated_at",
        ],
        [[
            "QA-0001",
            "Cepessa Sessions Dev",
            "/Users/ben/Developer/Cepessa Sessions/insights-worktree",
            BUILD,
            "desktop/build/Sessions Dev.app",
            "macOS 26, desktop accessory app",
            ENV,
            "local user of Sessions Dev",
            "Accessibility granted for agent-swift; Screen Recording not granted to screenshot helper (agent-swift screenshots worked)",
            "en-IL mixed Hebrew/English",
            "CEPESSA_INSIGHTS_USE_FIXTURE=1; cloud analysis off by default",
            "TypeSafe Jev optional; this run used fixture provider in the app",
            "Cloud analysis / decisions surface, settings, reader, export, sidecar, source, confirm/dismiss, RTL, prompt-package isolation",
            "Production Sessions.app; install; merge; push; live Jev on private recordings; recording/clips capture; transcription engine",
            "No production install, merge, publish, spend, or private transcript transmission",
            "Unverifiable (no user token budget)",
            "Unverifiable",
            "Baseline",
            "Inventory",
            NOW,
            NOW,
        ]],
    )

    # --- Features ---
    ws = wb.create_sheet("Features")
    features = [
        ["F-0001", "Settings", "Sessions Dev Settings > Cloud Analysis", "Local user",
         "As a user, I want cloud analysis off until I allow it, so that transcript text is not sent by surprise.",
         "Keep TypeSafe optional and explicit",
         "CloudAnalysisSettingsSection Toggle Allow TypeSafe analysis",
         f"{EV}/01-settings-cloud.png", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Not fully tested", "Toggle exists; footer states off by default"],
        ["F-0002", "Settings", "API key field", "Local user",
         "As a user, I want to save or remove a TypeSafe key in Keychain, so that analysis can run without embedding a key.",
         "Store credential on this Mac only",
         "LocalSessionInsightCredentialStore + Save/Remove key",
         f"{EV}/01-settings-cloud.png", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Not fully tested", "No key stored shown; live key not entered in Settings"],
        ["F-0003", "Reader", "Decisions block", "Local user",
         "As a user, I want a compact decisions area in the reader, so that I can review claims without leaving the transcript.",
         "Stay in the reading surface",
         "LocalSessionInsightsView in CepessaSessionReadingView",
         f"{EV}/02-reader-before-analyze.png", "Both", "Experimental", "Contract", "High", "Medium", "TRUE",
         "Not fully tested", ""],
        ["F-0004", "Reader", "Analyze with TypeSafe", "Local user",
         "As a user, I want to start analysis myself, so that I know when transcript text may be sent.",
         "Explicit start",
         "Analyze with TypeSafe button + disclosure copy",
         f"{EV}/02-reader-before-analyze.png", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Not fully tested", "Fixture env skips live send"],
        ["F-0005", "Reader", "Consent alert", "Local user",
         "As a user, I want a confirmation that names TypeSafe before a real send, so that I can refuse.",
         "Named disclosure",
         "alert Analyze with TypeSafe?",
         "Code only in this fixture environment", "Code only", "Experimental", "Contract", "Medium", "High", "TRUE",
         "Blocked", "Fixture env does not show consent"],
        ["F-0006", "Reader", "Proposal list", "Local user",
         "As a user, I want decisions, commitments, and open questions listed separately, so that I can tell them apart.",
         "Kind grouping",
         "itemGroup Decisions/Commitments/Open questions",
         f"{EV}/03-reader-after-analyze.png", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Not fully tested", ""],
        ["F-0007", "Reader", "Confirm", "Local user",
         "As a user, I want to confirm a proposal, so that my judgment is stored separately from the machine label.",
         "Human review",
         "reviewInsight confirmed",
         f"{EV}/06-after-confirm.png", "Both", "Experimental", "Contract", "High", "Medium", "TRUE",
         "Not fully tested", ""],
        ["F-0008", "Reader", "Dismiss", "Local user",
         "As a user, I want to dismiss a proposal, so that it leaves the working list without changing the transcript.",
         "Reject interpretation",
         "reviewInsight dismissed",
         f"{EV}/08-source-and-dismiss.png", "Both", "Experimental", "Contract", "High", "Medium", "TRUE",
         "Not fully tested", ""],
        ["F-0009", "Reader", "Source", "Local user",
         "As a user, I want to jump to the exact transcript span, so that I can check the quote.",
         "Evidence link",
         "revealInsightSource + highlightedTranscript",
         f"{EV}/08-source-and-dismiss.png", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Not fully tested", "Highlight weak; Source not a separate AX control"],
        ["F-0010", "Reader", "Export analysis", "Local user",
         "As a user, I want to export the analysis as Markdown, so that I can share provenance without fake deep links.",
         "Optional export",
         "Export analysis as Markdown",
         "AX button present; file dialog not completed", "Both", "Experimental", "Contract", "Medium", "Low", "TRUE",
         "Not fully tested", ""],
        ["F-0011", "Persistence", "insights.json sidecar", "Local user",
         "As a user, I want analysis to survive reopen without rewriting session.json, so that old recordings stay loadable.",
         "Safe persistence",
         "LocalSessionInsightStore sidecar",
         "insights.json on disk after confirm/dismiss", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Not fully tested", ""],
        ["F-0012", "Classification", "Suggestion vs decision", "Local user",
         "As a user, I want suggestions not listed as decisions, so that I do not treat maybe as a close.",
         "Do not task-ify discussion",
         "Fixture + UI: אולי נעלה בראשון absent from Decisions",
         f"{EV}/03-reader-after-analyze.png", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Not fully tested", "Fixture classification, not live Jev quality"],
        ["F-0013", "Reader", "Hebrew RTL", "Hebrew-speaking user",
         "As a Hebrew user, I want mixed transcript and quotes right-aligned, so that I can read them naturally.",
         "Readable Hebrew",
         "LocalTranscriptTextDirection + RTL rows",
         f"{EV}/02-reader-before-analyze.png", "Both", "Experimental", "Contract", "High", "Medium", "TRUE",
         "Not fully tested", ""],
        ["F-0014", "Reader", "Empty open questions", "Local user",
         "As a user, I want empty kinds to say none in the analyzed coverage, so that I do not think the whole meeting was scanned as questions.",
         "Honest empty copy",
         "None in the analyzed coverage.",
         f"{EV}/03-reader-after-analyze.png", "Both", "Experimental", "Contract", "High", "Low", "TRUE",
         "Not fully tested", ""],
        ["F-0015", "Export isolation", "Existing prompt package", "Local user",
         "As a user, I want ordinary session export packages unchanged, so that analysis is not silently injected.",
         "No silent export",
         "LocalSessionPromptPackageBuilder untouched for insights",
         "session-package.md has no TypeSafe/insights", "Both", "Experimental", "Contract", "High", "Medium", "TRUE",
         "Not fully tested", ""],
        ["F-0016", "Trust", "Experimental badge", "Local user",
         "As a user, I want the analysis labeled Experimental, so that I do not treat it as a guaranteed meeting record.",
         "No false confidence",
         "Experimental label",
         f"{EV}/02-reader-before-analyze.png", "Both", "Experimental", "Contract", "High", "Medium", "TRUE",
         "Not fully tested", ""],
        ["F-0017", "Neighbor", "Reader toolbar", "Local user",
         "As a user, I want Import, Transcribe, and transcript Export still in the toolbar, so that analysis does not replace existing tools.",
         "Preserve reader tools",
         "CepessaSessionReadingToolbar",
         f"{EV}/02-reader-before-analyze.png", "Both", "Experimental", "Inferred", "High", "Medium", "TRUE",
         "Not fully tested", "Presence only; transcribe not executed"],
        ["F-0018", "Analysis", "Cancel in-flight", "Local user",
         "As a user, I want to cancel a running analysis, so that I can stop a send.",
         "Stop in-flight work",
         "Cancel analysis button when status running",
         "Not observed; fixture finished instantly", "Code only", "Experimental", "Contract", "Medium", "Medium", "TRUE",
         "Blocked", "Too fast to click Cancel in this fixture"],
        ["F-0019", "Failure", "Missing key", "Local user",
         "As a user, I want a missing-key state that does not send the transcript, so that a misconfigured Mac stays local.",
         "Fail closed",
         "TypeSafeSessionInsightClient missingCredential",
         "Unit tests; Settings shows No key stored", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Not fully tested", "Live HTTP missing-key path not clicked in UI"],
        ["F-0020", "Accessibility", "Per-action controls", "Keyboard/VoiceOver user",
         "As a VoiceOver user, I want Source, Confirm, and Dismiss as separate controls, so that I can operate one action.",
         "Independent actions",
         "insightRow accessibilityElement children combine",
         "agent-swift: Source not found as element", "Both", "Experimental", "Contract", "High", "High", "TRUE",
         "Failing", "Combined AX node"],
    ]
    add_rows(
        ws,
        [
            "feature_id", "area", "surface", "actor", "user_story", "user_value",
            "code_evidence", "runtime_evidence", "source_presence", "availability",
            "expected_source", "confidence", "risk", "in_scope", "current_status", "notes",
        ],
        features,
    )

    # --- Behaviors ---
    ws = wb.create_sheet("Behaviors")
    behaviors = [
        ["B-0001", "F-0001", "Settings open", "Settings is open", "The user reads Cloud Analysis",
         "Copy says analysis is off by default and names TypeSafe", "Happy path", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0002", "F-0002", "Settings open, no key", "No key stored", "The user looks at the key row",
         "UI says No key stored and Save is disabled until text is entered", "Empty", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0003", "F-0003", "Insight fixture session open", "A saved Hebrew transcript is selected", "The reader loads",
         "A Decisions block appears above the transcript", "Happy path", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0004", "F-0004", "Reader open, unused analysis", "No analysis yet", "The user reads the helper text",
         "The disclosure names TypeSafe and says audio is not sent", "Happy path", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0005", "F-0004", "Reader open", "Analyze with TypeSafe is visible", "The user clicks Analyze",
         "The surface fills with proposals or an explicit empty/failure state", "Happy path", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0006", "F-0005", "Cloud path without fixture", "Feature enabled, no fixture flag", "The user clicks Analyze",
         "A confirmation names TypeSafe before network", "Permission", "Runtime UI", "FALSE", "High", "", "", "Fixture skips this"],
        ["B-0007", "F-0006", "After analyze", "Fixture analysis complete", "The user scans groups",
         "Decisions, Commitments, and Open questions are separate headings", "Happy path", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0008", "F-0006", "After analyze", "Fixture transcript includes an explicit decision", "The user looks at Decisions",
         "סיכמנו שעולים בראשון appears as a proposal", "Happy path", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0009", "F-0006", "After analyze", "Conditional span present", "The user looks at that row",
         "Conditional is visible on נעלה בראשון רק אם נטע תאשר", "Happy path", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0010", "F-0012", "After analyze", "Suggestion span present", "The user looks at Decisions",
         "אולי נעלה בראשון is not listed as a decision", "Negative", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0011", "F-0007", "Proposal visible", "An unreviewed decision is shown", "The user clicks Confirm",
         "The badge becomes Confirmed and Confirm disappears", "Happy path", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0012", "F-0008", "Proposal visible", "An unreviewed commitment is shown", "The user clicks Dismiss",
         "The badge becomes Dismissed and the transcript text is unchanged", "Happy path", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0013", "F-0009", "Proposal visible", "A proposal has a Source action", "The user activates Source",
         "The matching transcript span is selected and visibly highlighted", "Happy path", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0014", "F-0020", "VoiceOver/automation", "A proposal row is focused", "The user looks for Source as its own control",
         "Source, Confirm, and Dismiss are separate buttons", "Accessibility", "Runtime UI", "FALSE", "High", "", "", ""],
        ["B-0015", "F-0010", "After analysis", "Export analysis is enabled", "The user activates Export",
         "A save panel offers Markdown", "Happy path", "Runtime UI", "FALSE", "Low", "", "", ""],
        ["B-0016", "F-0011", "After confirm/dismiss", "insights.json exists", "The sidecar is read",
         "Review states match the UI (confirmed/dismissed)", "Recovery", "File + UI", "FALSE", "High", "", "", ""],
        ["B-0017", "F-0013", "Hebrew session", "Mixed Hebrew transcript", "The user reads quotes",
         "Hebrew blocks are right-aligned", "Happy path", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0018", "F-0014", "After analyze, no questions", "No open questions found", "The user reads Open questions",
         "Copy says none in the analyzed coverage", "Empty", "Runtime UI", "FALSE", "Low", "", "", ""],
        ["B-0019", "F-0015", "Session loaded", "Prompt package generated", "The package is inspected",
         "It does not mention TypeSafe or insights", "Negative", "File", "FALSE", "Medium", "", "", ""],
        ["B-0020", "F-0016", "Reader open", "Decisions block visible", "The user reads the heading",
         "Experimental is shown next to Decisions", "Happy path", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0021", "F-0017", "Reader open", "Toolbar visible", "The user looks at the toolbar",
         "Import Audio, Transcribe, and Export remain", "Happy path", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0022", "F-0018", "Analysis running", "Status is running", "The user looks for Cancel",
         "Cancel is available and stops the run", "Recovery", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0023", "F-0019", "Settings, empty key", "No key stored", "The user reads the key status",
         "No key stored is shown; Save stays disabled", "Empty", "Runtime UI", "FALSE", "Medium", "", "", ""],
        ["B-0024", "F-0006", "After analyze", "A commitment exists", "The user looks at Commitments",
         "אני אשלח לך מחר appears as a proposal", "Happy path", "Runtime UI", "FALSE", "High", "", "", ""],
    ]
    add_rows(
        ws,
        [
            "behavior_id", "feature_id", "preconditions", "given", "when", "then_expected",
            "behavior_type", "test_method", "automation_candidate", "risk",
            "current_result", "latest_test_run_id", "notes",
        ],
        behaviors,
    )

    # --- Test Runs (baseline) ---
    ws = wb.create_sheet("Test Runs")
    tests = [
        ["T-0001", "QA-0001", "B-0001", "Baseline", BUILD, ENV, "Dev app launched",
         "Opened Settings Dev; captured Cloud Analysis footer",
         "Off by default; TypeSafe named",
         "Footer and Cloud Analysis heading present; toggle was above the first crop but copy is on screen",
         "Pass", f"{EV}/01-settings-cloud.png", "grok-4.6", NOW, "Development fixture"],
        ["T-0002", "QA-0001", "B-0002", "Baseline", BUILD, ENV, "No key",
         "Read Save/Remove and status",
         "No key stored; Save disabled",
         "No key stored. Save key and Remove key disabled",
         "Pass", f"{EV}/01-settings-cloud.png", "grok-4.6", NOW, ""],
        ["T-0003", "QA-0001", "B-0003", "Baseline", BUILD, ENV, "Insight fixture session",
         "Opened session from status menu",
         "Decisions block above transcript",
         "Decisions Experimental and transcript visible",
         "Pass", f"{EV}/02-reader-before-analyze.png", "grok-4.6", NOW, "Development fixture"],
        ["T-0004", "QA-0001", "B-0004", "Baseline", BUILD, ENV, "No analysis",
         "Read helper text",
         "Names TypeSafe; audio not sent",
         "Cloud analysis is off until you start it... TypeSafe (Jev). Audio... not sent",
         "Pass", f"{EV}/02-reader-before-analyze.png", "grok-4.6", NOW, ""],
        ["T-0005", "QA-0001", "B-0005", "Baseline", BUILD, ENV, "Analyze visible",
         "Clicked Analyze transcript with TypeSafe",
         "Proposals appear",
         "Three proposals after click; status Analyzed 4 spans",
         "Pass", f"{EV}/03-reader-after-analyze.png", "grok-4.6", NOW, "Fixture provider, not live Jev"],
        ["T-0006", "QA-0001", "B-0006", "Baseline", BUILD, ENV, "Fixture flag on",
         "Clicked Analyze",
         "Consent alert before send",
         "No consent alert; fixture path runs immediately",
         "N/A", f"{EV}/03-reader-after-analyze.png", "grok-4.6", NOW, "Consent not exercised in this environment"],
        ["T-0007", "QA-0001", "B-0007", "Baseline", BUILD, ENV, "After analyze",
         "Read group headings",
         "Three groups",
         "Decisions, Commitments, Open questions all present",
         "Pass", f"{EV}/03-reader-after-analyze.png", "grok-4.6", NOW, ""],
        ["T-0008", "QA-0001", "B-0008", "Baseline", BUILD, ENV, "After analyze",
         "Read Decisions",
         "סיכמנו שעולים בראשון listed",
         "Listed as Proposal for Noam 0:16",
         "Pass", f"{EV}/03-reader-after-analyze.png", "grok-4.6", NOW, ""],
        ["T-0009", "QA-0001", "B-0009", "Baseline", BUILD, ENV, "After analyze",
         "Read second decision",
         "Conditional visible",
         "Proposal Conditional on נעלה בראשון רק אם נטע תאשר",
         "Pass", f"{EV}/03-reader-after-analyze.png", "grok-4.6", NOW, ""],
        ["T-0010", "QA-0001", "B-0010", "Baseline", BUILD, ENV, "After analyze",
         "Scan Decisions for אולי",
         "Suggestion not listed",
         "אולי only in transcript, not in Decisions",
         "Pass", f"{EV}/03-reader-after-analyze.png", "grok-4.6", NOW, ""],
        ["T-0011", "QA-0001", "B-0011", "Baseline", BUILD, ENV, "Unreviewed decision",
         "Clicked Confirm by coordinates",
         "Confirmed badge",
         "Green Confirmed; Confirm control gone; sidecar confirmed",
         "Pass", f"{EV}/06-after-confirm.png", "grok-4.6", NOW, "Clicked by coordinates because AX combined"],
        ["T-0012", "QA-0001", "B-0012", "Baseline", BUILD, ENV, "Unreviewed commitment",
         "Clicked Dismiss by coordinates",
         "Dismissed; transcript unchanged",
         "Dismissed badge; transcript still shows אני אשלח לך מחר",
         "Pass", f"{EV}/08-source-and-dismiss.png", "grok-4.6", NOW, ""],
        ["T-0013", "QA-0001", "B-0013", "Baseline", BUILD, ENV, "Source on first proposal",
         "Clicked Source; looked for highlight",
         "Exact span highlighted",
         "Source click registered; underline not clearly visible; all text already on screen so scroll proof is weak",
         "Fail", f"{EV}/08-source-and-dismiss.png", "grok-4.6", NOW, "Highlight too weak to prove"],
        ["T-0014", "QA-0001", "B-0014", "Baseline", BUILD, ENV, "After analyze",
         "agent-swift find text Source",
         "Separate Source/Confirm/Dismiss controls",
         "ELEMENT_NOT_FOUND for Source; row is one combined AX button",
         "Fail", "agent-swift ELEMENT_NOT_FOUND 9468d162", "grok-4.6", NOW, ""],
        ["T-0015", "QA-0001", "B-0015", "Baseline", BUILD, ENV, "After analyze",
         "Observed Export analysis as Markdown button",
         "Save panel",
         "Button present and enabled; save panel not completed to avoid extra files",
         "Blocked", "AX Export analysis as Markdown enabled", "grok-4.6", NOW, "Panel not driven"],
        ["T-0016", "QA-0001", "B-0016", "Baseline", BUILD, ENV, "After confirm/dismiss",
         "Read insights.json",
         "States match UI",
         "decision confirmed; commitment dismissed; conditional unreviewed",
         "Pass", "ui-root/.../insights.json", "grok-4.6", NOW, ""],
        ["T-0017", "QA-0001", "B-0017", "Baseline", BUILD, ENV, "Hebrew transcript",
         "Inspect alignment",
         "Hebrew right-aligned",
         "Transcript and quotes sit on the trailing edge",
         "Pass", f"{EV}/02-reader-before-analyze.png", "grok-4.6", NOW, ""],
        ["T-0018", "QA-0001", "B-0018", "Baseline", BUILD, ENV, "After analyze",
         "Read Open questions",
         "None in the analyzed coverage",
         "Exact copy present",
         "Pass", f"{EV}/03-reader-after-analyze.png", "grok-4.6", NOW, ""],
        ["T-0019", "QA-0001", "B-0019", "Baseline", BUILD, ENV, "Session load",
         "Grep session-package.md",
         "No TypeSafe/insights",
         "No matches",
         "Pass", "ui-root/.../Exports/session-package.md", "grok-4.6", NOW, ""],
        ["T-0020", "QA-0001", "B-0020", "Baseline", BUILD, ENV, "Reader",
         "Read heading",
         "Experimental visible",
         "Experimental next to Decisions",
         "Pass", f"{EV}/02-reader-before-analyze.png", "grok-4.6", NOW, ""],
        ["T-0021", "QA-0001", "B-0021", "Baseline", BUILD, ENV, "Reader toolbar",
         "Read toolbar",
         "Import Transcribe Export remain",
         "All three present",
         "Pass", f"{EV}/02-reader-before-analyze.png", "grok-4.6", NOW, ""],
        ["T-0022", "QA-0001", "B-0022", "Baseline", BUILD, ENV, "Fixture analyze",
         "Watch for running/Cancel",
         "Cancel available while running",
         "Analysis finished before Cancel could appear",
         "Blocked", "no running screenshot", "grok-4.6", NOW, ""],
        ["T-0023", "QA-0001", "B-0023", "Baseline", BUILD, ENV, "Settings",
         "Read key status",
         "No key stored",
         "No key stored shown",
         "Pass", f"{EV}/01-settings-cloud.png", "grok-4.6", NOW, ""],
        ["T-0024", "QA-0001", "B-0024", "Baseline", BUILD, ENV, "After analyze",
         "Read Commitments",
         "אני אשלח לך מחר listed",
         "Listed as Proposal",
         "Pass", f"{EV}/03-reader-after-analyze.png", "grok-4.6", NOW, ""],
    ]
    add_rows(
        ws,
        [
            "test_run_id", "run_id", "behavior_id", "phase", "commit_or_build", "environment",
            "preconditions", "steps", "expected", "actual", "result", "evidence",
            "tester", "tested_at", "notes",
        ],
        tests,
    )

    # --- Defects ---
    ws = wb.create_sheet("Defects")
    defects = [
        ["D-0001", "B-0014", "ax-row-combine-source-confirm-dismiss", "Accessibility", "High",
         "Source, Confirm, and Dismiss are one VoiceOver control",
         "Analyze a session; ask automation or VoiceOver for Source",
         "Each action is its own button",
         "find text Source returns ELEMENT_NOT_FOUND; row AXButton combines the whole proposal",
         f"{EV}/03-reader-after-analyze.png", "Reproduced", "Authorized by skill invocation",
         "insightRow uses accessibilityElement(children: .combine)", "", "", "", "",
         "", "Local reversible SwiftUI fix"],
        ["D-0002", "B-0013", "source-highlight-not-visible", "UX", "Medium",
         "Source does not make the quote obviously highlighted",
         "Click Source on סיכמנו שעולים בראשון",
         "The matching transcript span is clearly marked",
         "Click succeeded; underline is not visible in screenshots; whole transcript already on screen",
         f"{EV}/08-source-and-dismiss.png", "Reproduced", "Authorized by skill invocation",
         "highlightedTranscript uses underline only", "", "", "", "",
         "", "Contract asked for highlight"],
        ["D-0003", "B-0006", "consent-untested-in-fixture-env", "Integration/Environment", "High",
         "Live consent dialog not shown in this run",
         "Analyze without fixture flag",
         "TypeSafe confirmation before network",
         "Fixture environment skips consent",
         f"{EV}/03-reader-after-analyze.png", "Blocked", "Protected live send",
         "CEPESSA_INSIGHTS_USE_FIXTURE=1", "", "", "", "",
         "Needs a non-fixture run with an explicit send decision", "Not a product-code defect in this env"],
        ["D-0004", "B-0022", "cancel-not-observable", "Test harness", "Low",
         "Cancel never appeared because analysis finished immediately",
         "Start a slow analysis and click Cancel",
         "Cancel is visible while running",
         "Fixture completed in one frame",
         "no running screenshot", "Blocked", "n/a",
         "Fixture latency too low", "", "", "", "",
         "Needs delayed provider or live network", ""],
        ["D-0005", "B-0015", "export-panel-not-completed", "Test harness", "Low",
         "Markdown save panel was not completed",
         "Click Export analysis",
         "A markdown file is written",
         "Button enabled; panel not driven",
         "AX Export analysis as Markdown", "Blocked", "n/a",
         "Avoided extra modal files in this pass", "", "", "", "",
         "Can complete in retest", "Not a product defect"],
    ]
    add_rows(
        ws,
        [
            "defect_id", "behavior_ids", "fingerprint", "category", "severity", "title",
            "reproduction", "expected", "actual", "evidence", "status", "approval_state",
            "root_cause", "fix_reference", "fixed_build", "targeted_retest_id",
            "final_regression_id", "blocker_or_acceptance_reason", "notes",
        ],
        defects,
    )

    # --- UX Opportunities (filled after functional pass; placeholder rows allowed as proposed from baseline observations) ---
    ws = wb.create_sheet("UX Opportunities")
    add_rows(
        ws,
        [
            "opportunity_id", "feature_ids", "behavior_ids", "user_journey_step", "evidence_type",
            "evidence", "observed_friction", "proposed_change", "expected_user_benefit",
            "priority", "confidence", "effort", "risk", "validation_method", "visual_reference",
            "status", "notes",
        ],
        [
            ["UX-0001", "F-0009", "B-0013", "After clicking Source", "Screenshot",
             f"{EV}/08-source-and-dismiss.png",
             "The quote in the transcript does not stand out from neighboring lines",
             "Give the source span a lasting background mark and keep it on screen",
             "The user can check the quote without hunting",
             "P1", "High", "S", "Low", "Retest Source click with a visible mark in a screenshot",
             f"{EV}/08-source-and-dismiss.png", "Proposed",
             "If highlight remains a contract miss it is D-0002, not only an opportunity"],
            ["UX-0002", "F-0003", "B-0003", "First look at Decisions", "Screenshot",
             f"{EV}/02-reader-before-analyze.png",
             "Analyze with TypeSafe is a small caption-weight control on the trailing edge",
             "Keep the experimental label, but make the start action a regular trailing button",
             "The user can find the start action without scanning microcopy",
             "P2", "Medium", "S", "Low", "Ask a new user to start analysis without hints",
             f"{EV}/02-reader-before-analyze.png", "Proposed", "Opportunity, not a broken control"],
        ],
    )

    # --- Routing ---
    ws = wb.create_sheet("Routing")
    add_rows(
        ws,
        [
            "node_id", "role", "requested_model", "requested_effort", "effective_model",
            "effective_effort", "agent_identity", "runtime", "fallback_reason",
            "elapsed", "usage_credits", "acceptance", "rollout_ref",
        ],
        [
            ["N-0001", "orchestrator", "grok-4.6", "xhigh", "grok-4.6", "medium (host)",
             "Grok Build / T3 Code", "Grok harness",
             "Luna/Sol/Terra spawn pairs are not available; routing receipt Unverifiable for OpenAI tiers",
             "Unverifiable", "Unverifiable", "n/a", "Unverifiable"],
        ],
    )

    # --- UX coverage matrix ---
    ws = wb.create_sheet("UX Coverage")
    add_rows(
        ws,
        [
            "surface", "story", "evidence", "criteria", "ux_defects", "ux_opportunities",
            "no_finding", "limitation",
        ],
        [
            ["Settings Cloud Analysis", "F-0001/F-0002", f"{EV}/01-settings-cloud.png",
             "hierarchy, copy, empty key, destructive remove", "", "",
             "No extra opportunity beyond key empty state working", "Toggle cropped in first shot; footer visible"],
            ["Reader before analyze", "F-0003/F-0004/F-0016", f"{EV}/02-reader-before-analyze.png",
             "discoverability, trust copy, RTL", "", "UX-0002", "", ""],
            ["Reader after analyze", "F-0006/F-0012/F-0014", f"{EV}/03-reader-after-analyze.png",
             "grouping, badges, empty kind, RTL", "D-0001", "", "", ""],
            ["Confirm", "F-0007", f"{EV}/06-after-confirm.png",
             "feedback, badge change", "", "", "No evidence-backed opportunity found", ""],
            ["Dismiss", "F-0008", f"{EV}/08-source-and-dismiss.png",
             "feedback, transcript preserved", "", "", "No evidence-backed opportunity found", ""],
            ["Source", "F-0009", f"{EV}/08-source-and-dismiss.png",
             "orientation, highlight, scroll", "D-0002", "UX-0001", "", "Short transcript; scroll not proven"],
        ],
    )

    # --- Summary with formulas ---
    ws = wb.create_sheet("Summary")
    ws["A1"] = "metric"
    ws["B1"] = "value"
    style_header(ws, 2)
    metrics = [
        ("total_features", "=COUNTA(Features!A2:A50)"),
        ("in_scope_features", '=COUNTIF(Features!N2:N50,"TRUE")'),
        ("total_behaviors", "=COUNTA(Behaviors!A2:A80)"),
        ("baseline_pass", '=COUNTIFS(\'Test Runs\'!D2:D200,"Baseline",\'Test Runs\'!K2:K200,"Pass")'),
        ("baseline_fail", '=COUNTIFS(\'Test Runs\'!D2:D200,"Baseline",\'Test Runs\'!K2:K200,"Fail")'),
        ("baseline_blocked", '=COUNTIFS(\'Test Runs\'!D2:D200,"Baseline",\'Test Runs\'!K2:K200,"Blocked")'),
        ("baseline_na", '=COUNTIFS(\'Test Runs\'!D2:D200,"Baseline",\'Test Runs\'!K2:K200,"N/A")'),
        ("open_defects", '=COUNTIF(Defects!K2:K50,"Reproduced")+COUNTIF(Defects!K2:K50,"Open")'),
        ("blocked_defects", '=COUNTIF(Defects!K2:K50,"Blocked")'),
        ("verified_fixed", '=COUNTIF(Defects!K2:K50,"Verified fixed")'),
        ("ux_proposed", '=COUNTIF(\'UX Opportunities\'!P2:P30,"Proposed")'),
        ("run_status_manual_note", "Derived: Incomplete until D-0001/D-0002 fixed and final regression + verifier"),
        ("build", BUILD),
        ("updated_at", NOW),
    ]
    for i, (k, v) in enumerate(metrics, start=2):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
        ws.cell(i, 1).border = thin
        ws.cell(i, 2).border = thin
        ws.cell(i, 2).alignment = wrap
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 80

    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
